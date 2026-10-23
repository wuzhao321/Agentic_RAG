from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from kb.agent.tools.project_query import Project, ProjectArgs, query_projects
from kb.config import load_config
from kb.ingest.chunker import chunk_document
from kb.ingest.loaders_docx import load_docx
from kb.retrieval.bm25 import BM25Index


def run_demo() -> dict:
    with TemporaryDirectory(prefix="kb-demo-") as directory:
        path = Path(directory) / "合成接口说明.docx"
        document = Document()
        document.add_heading("CreateOrder 接口", level=1)
        document.add_paragraph("这是自动生成的测试资料，不是真实业务文档。")
        table = document.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ["字段", "类型", "说明"], strict=False):
            cell.text = value
        for values in [("order_id", "string", "订单编号"), ("timeout_ms", "integer", "3000 毫秒")]:
            for cell, value in zip(table.add_row().cells, values, strict=False):
                cell.text = value
        document.save(path)
        parsed = load_docx(path)
        chunks = chunk_document(
            parsed,
            {
                "doc_id": "synthetic-demo",
                "doc_name": path.name,
                "doc_type": "接口说明书",
                "project_name": "演示项目",
                "project_id": "DEMO",
                "source_format": "docx",
            },
            load_config("ingest"),
        )
        hits = BM25Index(chunks).search("timeout_ms", 3)
        projects = [
            Project(
                project_id="DEMO",
                project_name="演示项目",
                signed_date="2026-10-23",
                amount="100000",
                owner="测试负责人",
            )
        ]
        ledger = query_projects(projects, ProjectArgs(project_name="演示项目"), 80)
        return {
            "mode": "离线合成数据演示，仅验证解析/BM25/台账，不代表 BGE 或 LLM 效果",
            "chunks": len(chunks),
            "hits": [h.model_dump() for h in hits],
            "ledger": ledger,
        }
