import csv
import json

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from kb.ingest.chunker import chunk_document, table_chunks
from kb.ingest.indexer import ingest, read_chunks
from kb.ingest.loaders_docx import load_docx
from kb.ingest.loaders_pdf import load_pdf
from kb.ingest.metadata import scan
from kb.ingest.table_md import docx_table_markdown, rows_to_markdown
from kb.retrieval.vector import open_collection
from kb.schema import Block, ParsedDocument


def test_docx_code_indentation_is_preserved(tmp_path):
    doc = Document()
    doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("    return value", style="Source Code")
    path = tmp_path / "code.docx"
    doc.save(path)
    assert load_docx(path).blocks[0].text == "    return value"


def test_paragraph_context_preserved_without_crossing_boundaries(ingest_config):
    parsed = ParsedDocument(
        blocks=[
            Block(text="查询参数", section_path="查询", page=1),
            Block(text="不声明默认值则为必填。", section_path="查询", page=1),
            Block(
                text="| 字段 | 值 |\n| --- | --- |\n| needy | 必填 |",
                kind="table",
                section_path="查询",
                page=1,
            ),
            Block(text="表后补充", section_path="查询", page=1),
            Block(text="新一页", section_path="查询", page=2),
            Block(text="另一节", section_path="路径", page=2),
        ]
    )
    chunks = chunk_document(parsed, {"doc_id": "doc", "doc_name": "doc.pdf"}, ingest_config)
    assert len(chunks) == 5
    assert chunks[0].text == "查询参数\n\n不声明默认值则为必填。"
    assert chunks[1].metadata["chunk_type"] == "table"
    assert chunks[2].text == "表后补充" and chunks[3].metadata["page"] == 2
    assert chunks[4].metadata["section_path"] == "路径"
    assert parsed.blocks[0].text == "查询参数"


def test_docx_body_order_headings_and_merged_cells(tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("接口说明", 1)
    doc.add_paragraph("表格之前")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "合并内容"
    table.cell(2, 0).text = "timeout_ms"
    table.cell(2, 1).text = "3000 毫秒"
    doc.add_paragraph("表格之后")
    doc.save(path)
    parsed = load_docx(path)
    assert [b.kind for b in parsed.blocks] == ["text", "text", "table", "text"]
    assert parsed.blocks[2].text.count("合并内容") == 1
    assert "| timeout_ms | 3000 毫秒 |" in parsed.blocks[2].text
    assert all(b.section_path == "接口说明" for b in parsed.blocks)


def test_vertical_merge_and_nested_table():
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套文本"
    text, warnings = docx_table_markdown(table)
    assert text.count("纵向合并") == 1
    assert "嵌套文本" in text
    assert warnings


def test_table_escape_and_repeat_headers():
    text = rows_to_markdown(
        [["字段", "值"], ["a|b", "第一行\n第二行"]] + [[f"field_{i}", "x" * 25] for i in range(10)]
    )
    assert "a\\|b" in text and "第一行 第二行" in text
    chunks = table_chunks(text, "文档 > 章节 > 参数表", 110)
    assert len(chunks) > 1
    for chunk in chunks:
        assert "| 字段 | 值 |\n| --- | --- |" in chunk
    for row in text.splitlines()[2:]:
        assert sum(row in chunk.splitlines() for chunk in chunks) == 1


def test_long_table_row_is_not_truncated():
    value = "完整数值" * 100
    chunks = table_chunks(rows_to_markdown([["字段", "值"], ["x", value]]), "文档", 64)
    assert len(chunks) == 1 and value in chunks[0]


def test_chunk_ids_stable(tmp_path, ingest_config):
    path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("技术参数。" * 250)
    doc.save(path)
    parsed = load_docx(path)
    meta = {"doc_id": "id", "doc_name": "test.docx"}
    first = chunk_document(parsed, meta, ingest_config)
    second = chunk_document(parsed, meta, ingest_config)
    assert first == second
    assert len(first) > 1
    assert all(len(c.text) <= 512 for c in first)


def test_pdf_pages_and_empty_page_warning(tmp_path):
    path = tmp_path / "source.pdf"
    doc = pymupdf.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 30), "REPEATED HEADER")
        page.insert_text((72, 200), f"Page body {i + 1}")
    doc.new_page()
    doc.save(path)
    doc.close()
    parsed = load_pdf(path, 2, 0.6)
    assert parsed.page_count == 4
    assert {b.page for b in parsed.blocks} == {1, 2, 3}
    assert all("REPEATED HEADER" not in b.text for b in parsed.blocks)
    assert any("第 4 页无文本层" in warning for warning in parsed.warnings)


class TinyEmbedding:
    """Deterministic test double; never used as a production embedding."""

    def encode(self, texts):
        return [[1.0, len(text) / 100, 0.5] for text in texts]


def test_pdf_table_extraction_keeps_cells_and_avoids_duplicate_text(tmp_path):
    path = tmp_path / "table.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    for x in [70, 220, 370]:
        page.draw_line((x, 100), (x, 180))
    for y in [100, 140, 180]:
        page.draw_line((70, y), (370, y))
    for x, y, text in [
        (80, 125, "field"),
        (230, 125, "value"),
        (80, 165, "timeout_ms"),
        (230, 165, "3000 ms"),
    ]:
        page.insert_text((x, y), text)
    doc.save(path)
    doc.close()
    parsed = load_pdf(path, 2, 0.6)
    tables = [b for b in parsed.blocks if b.kind == "table"]
    assert len(tables) == 1
    assert "| timeout_ms | 3000 ms |" in tables[0].text
    assert sum(b.text.count("timeout_ms") for b in parsed.blocks) == 1


def prepare_registry(tmp_path, config):
    raw = tmp_path / "raw"
    raw.mkdir()
    doc = Document()
    doc.add_paragraph("CreateOrder 参数说明")
    path = raw / "interface.docx"
    doc.save(path)
    registry = tmp_path / "registry.csv"
    rows = scan(raw, registry, config)
    rows[0].update(doc_type="接口说明书", project_name="星河", project_id="P1")
    with registry.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)
    config = {
        **config,
        "raw_dir": str(raw),
        "registry": str(registry),
        "index_dir": str(tmp_path / "index"),
    }
    return path, config


def test_registry_preserves_manual_metadata(tmp_path, ingest_config):
    path, cfg = prepare_registry(tmp_path, ingest_config)
    rows = scan(path.parent, tmp_path / "registry.csv", cfg)
    assert rows[0]["project_name"] == "星河"
    assert rows[0]["doc_type"] == "接口说明书"


def test_chroma_generation_idempotency_and_update(tmp_path, ingest_config):
    path, cfg = prepare_registry(tmp_path, ingest_config)
    first = ingest(cfg, TinyEmbedding())
    active = tmp_path / "index" / "active.json"
    before = active.read_text()
    second = ingest(cfg, TinyEmbedding())
    assert first["total_chunks"] == second["total_chunks"] == 1
    assert second["unchanged"] and before == active.read_text()
    doc = Document()
    doc.add_paragraph("Updated DeleteOrder 接口")
    doc.save(path)
    ingest(cfg, TinyEmbedding())
    assert active.read_text() != before
    chunks = read_chunks(tmp_path / "index")
    assert len(chunks) == 1 and "DeleteOrder" in chunks[0].text
    manifest = json.loads(active.read_text())
    collection = open_collection(tmp_path / "index/chroma", manifest["collection"], cfg["chroma"])
    assert collection.count() == len(chunks)
    assert set(collection.get()["ids"]) == {c.id for c in chunks}


def test_failed_rebuild_keeps_active_generation(tmp_path, ingest_config):
    import pytest

    path, cfg = prepare_registry(tmp_path, ingest_config)
    ingest(cfg, TinyEmbedding())
    active = tmp_path / "index/active.json"
    before = active.read_text()
    doc = Document()
    doc.add_paragraph("Changed document")
    doc.save(path)

    class BrokenEmbedding:
        def encode(self, texts):
            raise RuntimeError("test: model failed")

    with pytest.raises(RuntimeError):
        ingest(cfg, BrokenEmbedding())
    assert active.read_text() == before
