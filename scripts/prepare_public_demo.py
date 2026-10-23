"""Reproducible public-data import. Does not touch formal tune/holdout datasets."""

import csv
import hashlib
import json
import re
import textwrap
from datetime import UTC, datetime
from pathlib import Path
from urllib.request import urlopen

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from kb.config import load_config
from kb.ingest.metadata import FIELDS, scan

REVISION = "69ddc2badf9da1123bba4d360225ae6b085985f6"
BASE = f"https://raw.githubusercontent.com/fastapi/fastapi/{REVISION}"
PAGES = {
    "body": "请求体",
    "query-params": "查询参数",
    "path-params": "路径参数",
    "response-model": "响应模型",
    "handling-errors": "错误处理",
}
CACHE = Path("data/public_sources")
RAW = Path("data/raw/public_demo")
OUT = Path("data/demo")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def download(url, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        with urlopen(url, timeout=60) as response:
            content = response.read()
        path.write_bytes(content)
    return path.read_text(encoding="utf-8")


def clean(text):
    text = re.sub(r"\{ #[^}]+\}", "", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "").replace("*", "").strip()


def make_docx(markdown, destination, source_url):
    doc = Document()
    doc.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            code = not code
            continue
        if not line.strip() or line.strip() == "///":
            continue
        if line.startswith("///"):
            line = line.split("|", 1)[-1]
        if not code and line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            doc.add_heading(clean(line.lstrip("#")), level=level)
        else:
            text = line if code else clean(line)
            if text:
                doc.add_paragraph(text, style="Source Code" if code else None)
    doc.add_heading("来源与转换说明", level=1)
    doc.add_paragraph(source_url)
    doc.add_paragraph(
        f"FastAPI 官方中文文档，MIT 许可，版本 {REVISION}。"
        "为知识库演示转换为 DOCX；保留文字和示例代码，省略图片。"
    )
    doc.save(destination)


def make_pdf(record, destination):
    """Render an explicitly labelled excerpt with a real text layer and ruled table."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((40, 40), record["title"], fontname="china-s", fontsize=16)
    page.insert_text(
        (40, 65), "公开公告摘录转换件，用于演示；不是原始合同。", fontname="china-s", fontsize=10
    )
    y = 90
    for key, value in {"字段": "公开信息", **record["fields"]}.items():
        wrapped = textwrap.wrap(value, width=32)
        height = max(28, len(wrapped) * 15 + 12)
        for x1, x2 in [(40, 140), (140, 555)]:
            page.draw_rect(pymupdf.Rect(x1, y, x2, y + height), color=(0.5, 0.5, 0.5))
        page.insert_text((48, y + 18), key, fontname="china-s", fontsize=9)
        for i, line in enumerate(wrapped):
            page.insert_text((148, y + 18 + i * 15), line, fontname="china-s", fontsize=10)
        y += height
    page.insert_text(
        (40, y + 25),
        "来源：官方公告；核对日期 2026-10-23；负责人未公开。",
        fontname="china-s",
        fontsize=10,
    )
    page.insert_text((40, y + 45), record["url"], fontsize=7)
    doc.save(destination)
    doc.close()


def main():
    for path in [CACHE, RAW, OUT]:
        path.mkdir(parents=True, exist_ok=True)
    manifest = []
    download(f"{BASE}/LICENSE", CACHE / "FASTAPI-LICENSE")
    (RAW / "LICENSE-FastAPI.txt").write_bytes((CACHE / "FASTAPI-LICENSE").read_bytes())
    for slug, title in PAGES.items():
        url = f"{BASE}/docs/zh/docs/tutorial/{slug}.md"
        source = CACHE / f"{slug}.md"
        markdown = download(url, source)
        code_sources = []

        def include(match, references=code_sources):
            relative = match.group(1)
            code = download(f"{BASE}/{relative}", CACHE / relative)
            references.append({"url": f"{BASE}/{relative}", "sha256": digest(CACHE / relative)})
            return f"\n```python\n{code}\n```\n"

        markdown = re.sub(r"\{\* ../../(docs_src/\S+\.py)[^}]*\*\}", include, markdown)
        target = RAW / f"FastAPI-{title}.docx"
        make_docx(markdown, target, url)
        manifest.append(
            {
                "id": slug,
                "title": f"FastAPI · {title}",
                "file_name": f"public_demo/{target.name}",
                "source_url": url,
                "source_path": str(source),
                "source_sha256": digest(source),
                "artifact_sha256": digest(target),
                "revision": REVISION,
                "code_sources": code_sources,
                "license": "MIT",
                "transformation": "Markdown 转 DOCX；展开固定版本示例代码；省略图片",
                "doc_type": "接口说明书",
                "project_id": "FASTAPI",
                "project_name": "FastAPI",
            }
        )
    records = json.loads(Path("docs/public-data/procurement-snapshots.json").read_text())
    for i, record in enumerate(records):
        for value in record["fields"].values():
            if value not in record["snapshot"]:
                raise ValueError(f"公告快照未包含字段：{value}")
        suffix = ".pdf" if i == 0 else ".docx"
        target = RAW / (record["title"] + suffix)
        if suffix == ".pdf":
            make_pdf(record, target)
        else:
            doc = Document()
            doc.add_heading(record["title"], 0)
            doc.add_paragraph("公开公告摘录转换件，用于演示；不是原始合同。")
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for k, v in {"字段": "公开信息", **record["fields"]}.items():
                cells = table.add_row().cells
                cells[0].text, cells[1].text = k, v
            doc.add_paragraph("来源：" + record["url"])
            doc.add_paragraph("核对日期：2026-10-23；项目负责人未公开。")
            doc.save(target)
        manifest.append(
            {
                "id": record["id"],
                "title": record["title"],
                "file_name": f"public_demo/{target.name}",
                "source_url": record["url"],
                "source_path": "docs/public-data/procurement-snapshots.json",
                "source_sha256": hashlib.sha256(record["snapshot"].encode()).hexdigest(),
                "artifact_sha256": digest(target),
                "doc_type": "其他",
                "transformation": "官方网页检索快照中的已核对字段转表格；省略联系方式；非原始合同",
                "project_id": record["ledger"]["project_id"],
                "project_name": record["ledger"]["project_name"],
            }
        )
    # Preserve unrelated existing ledger rows; only replace these two public records.
    ledger_path = Path("data/ledger/projects.json")
    existing = json.loads(ledger_path.read_text()) if ledger_path.exists() else []
    public_ids = {r["ledger"]["project_id"] for r in records}
    existing = [r for r in existing if r["project_id"] not in public_ids]
    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    ledger_path.write_text(
        json.dumps(existing + [r["ledger"] for r in records], ensure_ascii=False, indent=2)
    )
    config = load_config("ingest")
    rows = scan(Path(config["raw_dir"]), Path(config["registry"]), config)
    metadata = {m["file_name"]: m for m in manifest}
    for row in rows:
        if row["file_name"] in metadata:
            item = metadata[row["file_name"]]
            row.update({key: item[key] for key in ["doc_type", "project_id", "project_name"]})
    with Path(config["registry"]).open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    result = {
        "created_at": datetime.now(UTC).isoformat(),
        "demo_only": True,
        "human_reviewed": False,
        "review_status": "assistant_source_checked",
        "documents": manifest,
    }
    (OUT / "manifest.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(
        json.dumps(
            {
                "documents": len(manifest),
                "ledger_rows": len(records),
                "manifest": str(OUT / "manifest.json"),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
