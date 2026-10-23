"""Exercise actual BGE + Chroma on synthetic DOCX, without a cloud API or business data."""

import csv
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from kb.config import load_config
from kb.ingest.indexer import ingest, read_chunks
from kb.ingest.metadata import scan
from kb.retrieval.bm25 import BM25Index
from kb.retrieval.pipeline import RetrievalPipeline
from kb.retrieval.reranker import Reranker
from kb.retrieval.vector import LocalEmbedding, VectorIndex, open_collection


def main():
    ingest_cfg, retrieval_cfg = load_config("ingest"), load_config("retrieval")
    with TemporaryDirectory(prefix="kb-real-model-smoke-") as directory:
        root = Path(directory)
        raw = root / "raw"
        raw.mkdir()
        document = Document()
        document.add_heading("CreateOrder 接口参数", 1)
        document.add_paragraph("以下为合成测试数据。")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(
            table.rows, [("字段", "默认值"), ("timeout_ms", "3000 毫秒")], strict=True
        ):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        document.save(raw / "合成接口.docx")
        config = {
            **ingest_cfg,
            "raw_dir": str(raw),
            "registry": str(root / "registry.csv"),
            "index_dir": str(root / "index"),
        }
        rows = scan(raw, Path(config["registry"]), config)
        rows[0].update(doc_type="接口说明书", project_name="演示项目", project_id="DEMO")
        with Path(config["registry"]).open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)
        embedding = LocalEmbedding(config["embedding"])
        first = ingest(config, embedding)
        second = ingest(config, embedding)
        assert first["total_chunks"] == second["total_chunks"] and second["unchanged"]
        index = Path(config["index_dir"])
        manifest = json.loads((index / "active.json").read_text())
        collection = open_collection(index / "chroma", manifest["collection"], config["chroma"])
        chunks = read_chunks(index)
        pipeline = RetrievalPipeline(
            VectorIndex(collection, embedding, retrieval_cfg),
            BM25Index(chunks),
            Reranker(retrieval_cfg["reranker"]),
            retrieval_cfg,
            ["演示项目"],
        )
        hits = pipeline.search("演示项目的 CreateOrder 接口 timeout_ms 默认值是多少？")
        assert hits and "3000 毫秒" in hits[0].chunk.text
        no_hits = pipeline.search("火星上有多少只企鹅？")
        assert not no_hits
        print(
            json.dumps(
                {
                    "synthetic_data": True,
                    "real_bge_models": True,
                    "cloud_calls": 0,
                    "chunks": len(chunks),
                    "idempotent": True,
                    "matched_source": hits[0].chunk.metadata,
                    "matched_text": hits[0].chunk.text,
                    "irrelevant_query_hits": len(no_hits),
                },
                ensure_ascii=False,
                indent=2,
            )
        )


if __name__ == "__main__":
    main()
